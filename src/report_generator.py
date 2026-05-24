"""
Report Generator
Creates PDF and DOCX reports with cost analysis and recommendations
"""

from datetime import datetime
from typing import Dict, List
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image
from reportlab.lib import colors
from reportlab.graphics.shapes import Drawing, Rect, String
from reportlab.graphics.charts.piecharts import Pie
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.charts.linecharts import HorizontalLineChart
from reportlab.graphics import renderPDF
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io
import tempfile
import os

class ReportGenerator:
    def __init__(self):
        """Initialize report generator"""
        self.severity_colors = {
            'Critical': colors.red,
            'High': colors.orange,
            'Medium': colors.yellow,
            'Low': colors.lightgreen
        }
    
    def generate_pdf(
        self,
        cost_data: Dict,
        ri_data: Dict,
        recommendations: List[Dict],
        output_path: str
    ):
        """Generate PDF report"""
        doc = SimpleDocTemplate(output_path, pagesize=letter)
        story = []
        styles = getSampleStyleSheet()
        
        # Title with CloudXcelAI Branding
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=28,
            textColor=colors.HexColor('#0066cc'),
            spaceAfter=10,
            alignment=1  # Center
        )
        subtitle_style = ParagraphStyle(
            'Subtitle',
            parent=styles['Normal'],
            fontSize=14,
            textColor=colors.HexColor('#003d7a'),
            spaceAfter=30,
            alignment=1  # Center
        )
        
        story.append(Paragraph("CLOUDXCELAI FINOPTIMIZER", title_style))
        story.append(Paragraph("Cloud Financial Operations & Optimization Report", subtitle_style))
        story.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", styles['Normal']))
        story.append(Spacer(1, 0.3*inch))
        
        # Executive Summary
        story.append(Paragraph("Executive Summary", styles['Heading2']))
        summary_data = [
            ['Metric', 'Value'],
            ['Total Cost (6 months)', f"${cost_data['total_cost']:,.2f}"],
            ['Monthly Average', f"${cost_data['total_cost']/6:,.2f}"],
            ['Total Recommendations', str(len(recommendations))],
            ['Critical/High Priority', str(sum(1 for r in recommendations if r.get('severity') in ['Critical', 'High']))]
        ]
        
        summary_table = Table(summary_data, colWidths=[3*inch, 2*inch])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        story.append(summary_table)
        story.append(Spacer(1, 0.3*inch))
        
        # RI Analysis
        story.append(Paragraph("Reserved Instance Analysis", styles['Heading2']))
        ri_recommendations = ri_data['recommendations'].get('Recommendations', [])
        story.append(Paragraph(f"Active RIs: {len(ri_data['current_ris'])}", styles['Normal']))
        story.append(Paragraph(f"RI Purchase Recommendations: {len(ri_recommendations)}", styles['Normal']))
        story.append(Spacer(1, 0.2*inch))
        
        # Recommendations
        story.append(PageBreak())
        story.append(Paragraph("Cost Optimization Recommendations", styles['Heading2']))
        story.append(Spacer(1, 0.2*inch))
        
        for idx, rec in enumerate(recommendations, 1):
            # Recommendation header
            rec_title = f"{idx}. {rec.get('title', 'Recommendation')}"
            story.append(Paragraph(rec_title, styles['Heading3']))
            
            # Details table
            rec_data = [
                ['Severity', rec.get('severity', 'N/A')],
                ['Timeline', rec.get('timeline', 'N/A')],
                ['Estimated Savings', rec.get('estimated_savings', 'N/A')],
                ['Impact', rec.get('impact', 'N/A')]
            ]
            
            rec_table = Table(rec_data, colWidths=[1.5*inch, 4*inch])
            rec_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (0, -1), colors.lightgrey),
                ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('FONTSIZE', (0, 0), (-1, -1), 10)
            ]))
            story.append(rec_table)
            story.append(Spacer(1, 0.1*inch))
            
            # Description
            story.append(Paragraph(f"<b>Description:</b> {rec.get('description', 'N/A')}", styles['Normal']))
            story.append(Spacer(1, 0.1*inch))
            
            # Implementation steps
            if 'steps' in rec and rec['steps']:
                story.append(Paragraph("<b>Implementation Steps:</b>", styles['Normal']))
                for step in rec['steps']:
                    story.append(Paragraph(f"• {step}", styles['Normal']))
            
            story.append(Spacer(1, 0.3*inch))
        
        # Add Cost Analysis Charts
        story.append(PageBreak())
        story.append(Paragraph("Cost Analysis & Trends", styles['Heading2']))
        story.append(Spacer(1, 0.2*inch))
        
        # Cost breakdown chart
        story.append(self._create_cost_breakdown_chart(cost_data))
        story.append(Spacer(1, 0.3*inch))
        
        # Recommendations by severity chart
        story.append(Paragraph("Recommendations by Severity", styles['Heading3']))
        story.append(self._create_recommendations_chart(recommendations))
        story.append(Spacer(1, 0.3*inch))
        
        # Build PDF
        doc.build(story)
    
    def _create_cost_breakdown_chart(self, cost_data):
        """Create a pie chart for cost breakdown"""
        drawing = Drawing(400, 200)
        
        # Create pie chart
        pie = Pie()
        pie.x = 100
        pie.y = 50
        pie.width = 150
        pie.height = 150
        
        # Sample data (in real implementation, this would come from actual cost data)
        pie.data = [40, 25, 20, 15]  # EC2, S3, RDS, Other
        pie.labels = ['EC2', 'S3', 'RDS', 'Other']
        pie.slices.strokeWidth = 0.5
        
        # Colors
        pie.slices[0].fillColor = colors.HexColor('#0066cc')
        pie.slices[1].fillColor = colors.HexColor('#28a745')
        pie.slices[2].fillColor = colors.HexColor('#ffc107')
        pie.slices[3].fillColor = colors.HexColor('#dc3545')
        
        drawing.add(pie)
        
        # Title
        title = String(200, 180, 'Cost Distribution by Service', textAnchor='middle')
        title.fontName = 'Helvetica-Bold'
        title.fontSize = 12
        drawing.add(title)
        
        return drawing
    
    def _create_recommendations_chart(self, recommendations):
        """Create a bar chart for recommendations by severity"""
        drawing = Drawing(400, 200)
        
        # Count recommendations by severity
        severity_counts = {'Critical': 0, 'High': 0, 'Medium': 0, 'Low': 0}
        for rec in recommendations:
            severity = rec.get('severity', 'Medium')
            if severity in severity_counts:
                severity_counts[severity] += 1
        
        # Create bar chart
        chart = VerticalBarChart()
        chart.x = 50
        chart.y = 50
        chart.height = 125
        chart.width = 300
        
        # Data
        chart.data = [[severity_counts['Critical'], severity_counts['High'], 
                      severity_counts['Medium'], severity_counts['Low']]]
        chart.categoryAxis.categoryNames = ['Critical', 'High', 'Medium', 'Low']
        
        # Styling
        chart.bars[0].fillColor = colors.HexColor('#dc3545')
        chart.valueAxis.valueMin = 0
        chart.categoryAxis.labels.boxAnchor = 'ne'
        chart.categoryAxis.labels.dx = 8
        chart.categoryAxis.labels.dy = -2
        
        drawing.add(chart)
        
        # Title
        title = String(200, 180, 'Recommendations by Severity Level', textAnchor='middle')
        title.fontName = 'Helvetica-Bold'
        title.fontSize = 12
        drawing.add(title)
        
        return drawing
    
    def generate_pdf_with_charts(
        self,
        cost_data: Dict,
        ri_data: Dict,
        recommendations: List[Dict],
        all_data: Dict,
        output_path: str
    ):
        """Generate comprehensive PDF report with charts and visualizations"""
        doc = SimpleDocTemplate(output_path, pagesize=letter)
        story = []
        styles = getSampleStyleSheet()
        
        # Custom styles
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=28,
            textColor=colors.HexColor('#0066cc'),
            spaceAfter=10,
            alignment=1
        )
        subtitle_style = ParagraphStyle(
            'Subtitle',
            parent=styles['Normal'],
            fontSize=14,
            textColor=colors.HexColor('#003d7a'),
            spaceAfter=30,
            alignment=1
        )
        
        # Title Page
        story.append(Paragraph("CLOUDXCELAI FINOPTIMIZER", title_style))
        story.append(Paragraph("Comprehensive Cloud Financial Operations Report", subtitle_style))
        story.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", styles['Normal']))
        story.append(Spacer(1, 0.5*inch))
        
        # Executive Summary
        story.append(Paragraph("Executive Summary", styles['Heading2']))
        
        total_cost = cost_data.get('total_cost', 0)
        monthly_avg = total_cost / 6 if total_cost else 0
        idle_waste = all_data.get('idle_resources', {}).get('total_monthly_waste', 0)
        
        summary_data = [
            ['Metric', 'Value'],
            ['Total Cost (6 months)', f"${total_cost:,.2f}"],
            ['Monthly Average', f"${monthly_avg:,.2f}"],
            ['Total Recommendations', str(len(recommendations))],
            ['Critical/High Priority', str(sum(1 for r in recommendations if r.get('severity') in ['Critical', 'High']))],
            ['Monthly Savings Potential', f"${idle_waste:,.2f}"],
            ['Annual Savings Potential', f"${idle_waste * 12:,.2f}"]
        ]
        
        summary_table = Table(summary_data, colWidths=[3*inch, 2*inch])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0066cc')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        story.append(summary_table)
        story.append(Spacer(1, 0.3*inch))
        
        # Cost Analysis with Charts
        story.append(PageBreak())
        story.append(Paragraph("Cost Analysis & Visualization", styles['Heading2']))
        story.append(Spacer(1, 0.2*inch))
        
        # Cost breakdown chart
        story.append(self._create_cost_breakdown_chart(cost_data))
        story.append(Spacer(1, 0.3*inch))
        
        # Resource distribution chart
        story.append(Paragraph("Resource Distribution", styles['Heading3']))
        story.append(self._create_resource_distribution_chart(all_data))
        story.append(Spacer(1, 0.3*inch))
        
        # Recommendations Analysis
        story.append(PageBreak())
        story.append(Paragraph("Recommendations Analysis", styles['Heading2']))
        story.append(Spacer(1, 0.2*inch))
        
        # Recommendations by severity chart
        story.append(self._create_recommendations_chart(recommendations))
        story.append(Spacer(1, 0.3*inch))
        
        # Detailed Recommendations
        story.append(Paragraph("Detailed Recommendations", styles['Heading3']))
        story.append(Spacer(1, 0.2*inch))
        
        for idx, rec in enumerate(recommendations[:10], 1):  # Limit to top 10
            rec_title = f"{idx}. {rec.get('title', 'Recommendation')}"
            story.append(Paragraph(rec_title, styles['Heading4']))
            
            rec_data = [
                ['Severity', rec.get('severity', 'N/A')],
                ['Timeline', rec.get('timeline', 'N/A')],
                ['Estimated Savings', rec.get('estimated_savings', 'N/A')],
                ['Impact', rec.get('impact', 'N/A')]
            ]
            
            rec_table = Table(rec_data, colWidths=[1.5*inch, 4*inch])
            rec_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (0, -1), colors.lightgrey),
                ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('FONTSIZE', (0, 0), (-1, -1), 10)
            ]))
            story.append(rec_table)
            story.append(Spacer(1, 0.2*inch))
        
        # RI Analysis
        story.append(PageBreak())
        story.append(Paragraph("Reserved Instance Analysis", styles['Heading2']))
        ri_recommendations = ri_data.get('recommendations', {}).get('Recommendations', [])
        story.append(Paragraph(f"Active RIs: {len(ri_data.get('current_ris', []))}", styles['Normal']))
        story.append(Paragraph(f"RI Purchase Recommendations: {len(ri_recommendations)}", styles['Normal']))
        story.append(Spacer(1, 0.2*inch))
        
        # Build PDF
        doc.build(story)
    
    def _create_resource_distribution_chart(self, all_data):
        """Create a bar chart for resource distribution"""
        drawing = Drawing(400, 200)
        
        # Get resource counts
        ec2_count = len(all_data.get('rightsizing', {}).get('recommendations', []))
        s3_count = all_data.get('s3_data', {}).get('total_buckets', 0)
        rds_count = all_data.get('rds_data', {}).get('total_instances', 0)
        lambda_count = all_data.get('lambda_data', {}).get('total_functions', 0)
        
        # Create bar chart
        chart = VerticalBarChart()
        chart.x = 50
        chart.y = 50
        chart.height = 125
        chart.width = 300
        
        # Data
        chart.data = [[ec2_count, s3_count, rds_count, lambda_count]]
        chart.categoryAxis.categoryNames = ['EC2', 'S3', 'RDS', 'Lambda']
        
        # Styling
        chart.bars[0].fillColor = colors.HexColor('#0066cc')
        chart.valueAxis.valueMin = 0
        chart.categoryAxis.labels.boxAnchor = 'ne'
        chart.categoryAxis.labels.dx = 8
        chart.categoryAxis.labels.dy = -2
        
        drawing.add(chart)
        
        # Title
        title = String(200, 180, 'Resource Count by Service', textAnchor='middle')
        title.fontName = 'Helvetica-Bold'
        title.fontSize = 12
        drawing.add(title)
        
        return drawing
    
    def generate_kpi_pdf(
        self,
        maturity_data: Dict,
        cost_data: Dict,
        output_path: str
    ):
        """Generate KPI-focused PDF with charts and visualizations"""
        doc = SimpleDocTemplate(output_path, pagesize=letter)
        story = []
        styles = getSampleStyleSheet()
        
        # Custom styles
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=28,
            textColor=colors.HexColor('#0066cc'),
            spaceAfter=10,
            alignment=1
        )
        subtitle_style = ParagraphStyle(
            'Subtitle',
            parent=styles['Normal'],
            fontSize=14,
            textColor=colors.HexColor('#003d7a'),
            spaceAfter=30,
            alignment=1
        )
        kpi_header_style = ParagraphStyle(
            'KPIHeader',
            parent=styles['Heading2'],
            fontSize=18,
            textColor=colors.HexColor('#0066cc'),
            spaceAfter=15
        )
        
        # Title Page
        story.append(Paragraph("CLOUDXCELAI FINOPTIMIZER", title_style))
        story.append(Paragraph("FinOps Maturity Assessment Report", subtitle_style))
        story.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", styles['Normal']))
        story.append(Spacer(1, 0.5*inch))
        
        # Executive Summary with Overall Score
        story.append(Paragraph("Executive Summary", kpi_header_style))
        
        # Overall maturity score card
        overall_score = maturity_data.get('overall_maturity', 0)
        maturity_level = maturity_data.get('maturity_level', 'Unknown')
        
        # Create a visual score card
        score_data = [
            ['Overall FinOps Maturity Score', f"{overall_score:.1f}/4.0"],
            ['Maturity Level', maturity_level],
            ['Assessment Date', datetime.now().strftime('%Y-%m-%d')],
            ['Potential Monthly Savings', f"${maturity_data.get('potential_monthly_savings', 0):,.2f}"]
        ]
        
        score_table = Table(score_data, colWidths=[3*inch, 2.5*inch])
        score_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0066cc')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 14),
            ('FONTSIZE', (0, 1), (-1, -1), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE')
        ]))
        story.append(score_table)
        story.append(Spacer(1, 0.3*inch))
        
        # Maturity Level Description
        level_descriptions = {
            'Crawl': 'Ad-hoc processes, limited visibility, reactive cost control',
            'Walk': 'Defined processes, partial automation, beginning operational discipline',
            'Run': 'Mature governance, fully automated, cost-efficient systems',
            'Optimize / Innovate': 'Continuous optimization, predictive analytics, business value tracking'
        }
        
        description = level_descriptions.get(maturity_level, 'Assessment in progress')
        story.append(Paragraph(f"<b>Maturity Description:</b> {description}", styles['Normal']))
        story.append(Spacer(1, 0.4*inch))
        
        # KPI Dimensions Analysis
        story.append(PageBreak())
        story.append(Paragraph("KPI Dimensions Analysis", kpi_header_style))
        
        # Create charts for each KPI dimension
        dimensions = [
            ('Tagged Resources', maturity_data.get('tag_score', 0), maturity_data.get('tagged_percentage', 0)),
            ('Budget Utilization', maturity_data.get('budget_score', 0), maturity_data.get('budget_utilization', 0)),
            ('Savings Optimization', maturity_data.get('savings_score', 0), maturity_data.get('savings_percentage', 0)),
            ('Resource Utilization', maturity_data.get('utilization_score', 0), maturity_data.get('utilization_percentage', 0))
        ]
        
        # KPI Summary Table
        kpi_data = [['KPI Dimension', 'Score (0-4)', 'Percentage', 'Status']]
        
        for dim_name, score, percentage in dimensions:
            if score >= 3:
                status = '✓ Good'
            elif score >= 2:
                status = '⚠ Fair'
            else:
                status = '✗ Needs Improvement'
            
            kpi_data.append([
                dim_name,
                f"{score}/4",
                f"{percentage:.1f}%",
                status
            ])
        
        kpi_table = Table(kpi_data, colWidths=[2*inch, 1*inch, 1*inch, 1.5*inch])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0066cc')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('FONTSIZE', (0, 1), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE')
        ]))
        story.append(kpi_table)
        story.append(Spacer(1, 0.3*inch))
        
        # Create visual charts for KPIs
        story.append(self._create_maturity_chart(dimensions))
        story.append(Spacer(1, 0.3*inch))
        
        # Detailed KPI Analysis
        story.append(PageBreak())
        story.append(Paragraph("Detailed KPI Analysis", kpi_header_style))
        
        for dim_name, score, percentage in dimensions:
            story.append(Paragraph(f"{dim_name}", styles['Heading3']))
            
            # Score visualization
            story.append(self._create_score_bar(dim_name, score))
            
            # Analysis text
            analysis_text = self._get_kpi_analysis(dim_name, score, percentage)
            story.append(Paragraph(analysis_text, styles['Normal']))
            story.append(Spacer(1, 0.2*inch))
        
        # Cost Analysis Section
        story.append(PageBreak())
        story.append(Paragraph("Cost Analysis", kpi_header_style))
        
        # Cost metrics
        total_cost = cost_data.get('total_cost', 0)
        monthly_avg = total_cost / 6 if total_cost else 0
        potential_savings = maturity_data.get('potential_monthly_savings', 0)
        
        cost_data_table = [
            ['Cost Metric', 'Value'],
            ['Total Cost (6 months)', f"${total_cost:,.2f}"],
            ['Monthly Average', f"${monthly_avg:,.2f}"],
            ['Potential Monthly Savings', f"${potential_savings:,.2f}"],
            ['Annual Savings Potential', f"${potential_savings * 12:,.2f}"]
        ]
        
        cost_table = Table(cost_data_table, colWidths=[3*inch, 2*inch])
        cost_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#28a745')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.lightgreen),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        story.append(cost_table)
        story.append(Spacer(1, 0.3*inch))
        
        # Recommendations
        story.append(Paragraph("Recommendations", kpi_header_style))
        
        recommendations = self._generate_kpi_recommendations(maturity_data)
        for idx, rec in enumerate(recommendations, 1):
            story.append(Paragraph(f"{idx}. {rec}", styles['Normal']))
            story.append(Spacer(1, 0.1*inch))
        
        # Build PDF
        doc.build(story)
    
    def _create_maturity_chart(self, dimensions):
        """Create a bar chart for maturity dimensions"""
        drawing = Drawing(400, 200)
        
        # Create bar chart
        chart = VerticalBarChart()
        chart.x = 50
        chart.y = 50
        chart.height = 125
        chart.width = 300
        
        # Data
        chart.data = [[dim[1] for dim in dimensions]]  # Scores
        chart.categoryAxis.categoryNames = [dim[0] for dim in dimensions]
        
        # Styling
        chart.bars[0].fillColor = colors.HexColor('#0066cc')
        chart.valueAxis.valueMin = 0
        chart.valueAxis.valueMax = 4
        chart.categoryAxis.labels.boxAnchor = 'ne'
        chart.categoryAxis.labels.dx = 8
        chart.categoryAxis.labels.dy = -2
        chart.categoryAxis.labels.angle = 30
        
        drawing.add(chart)
        
        # Title
        title = String(200, 180, 'FinOps Maturity Scores by Dimension', textAnchor='middle')
        title.fontName = 'Helvetica-Bold'
        title.fontSize = 12
        drawing.add(title)
        
        return drawing
    
    def _create_score_bar(self, dimension_name, score):
        """Create a horizontal score bar for individual KPI"""
        drawing = Drawing(400, 60)
        
        # Background bar
        bg_bar = Rect(50, 25, 300, 20)
        bg_bar.fillColor = colors.lightgrey
        bg_bar.strokeColor = colors.black
        drawing.add(bg_bar)
        
        # Score bar
        score_width = (score / 4) * 300
        score_bar = Rect(50, 25, score_width, 20)
        
        # Color based on score
        if score >= 3:
            score_bar.fillColor = colors.green
        elif score >= 2:
            score_bar.fillColor = colors.yellow
        else:
            score_bar.fillColor = colors.red
        
        drawing.add(score_bar)
        
        # Score text
        score_text = String(55, 30, f"{score:.1f}/4.0", textAnchor='start')
        score_text.fontName = 'Helvetica-Bold'
        score_text.fontSize = 10
        drawing.add(score_text)
        
        return drawing
    
    def _get_kpi_analysis(self, dimension_name, score, percentage):
        """Generate analysis text for each KPI dimension"""
        analyses = {
            'Tagged Resources': f"Resource tagging coverage is at {percentage:.1f}%. " +
                              ("Excellent tagging discipline supports cost allocation and governance." if score >= 3 else
                               "Improve tagging strategy to enhance cost visibility and accountability." if score >= 2 else
                               "Critical: Implement comprehensive tagging strategy immediately."),
            
            'Budget Utilization': f"Budget utilization is at {percentage:.1f}%. " +
                                ("Optimal budget management with good forecasting accuracy." if score >= 3 else
                                 "Budget management needs refinement for better cost control." if score >= 2 else
                                 "Poor budget utilization indicates need for better planning and monitoring."),
            
            'Savings Optimization': f"Potential savings identified at {percentage:.1f}% of current spend. " +
                                  ("Well-optimized environment with minimal waste." if score >= 3 else
                                   "Moderate optimization opportunities available." if score >= 2 else
                                   "Significant optimization potential - immediate action recommended."),
            
            'Resource Utilization': f"Resource utilization efficiency is at {percentage:.1f}%. " +
                                  ("Excellent resource efficiency with minimal waste." if score >= 3 else
                                   "Good utilization with some improvement opportunities." if score >= 2 else
                                   "Poor utilization - review rightsizing and idle resource cleanup.")
        }
        
        return analyses.get(dimension_name, f"Analysis for {dimension_name}: Score {score:.1f}/4.0")
    
    def _generate_kpi_recommendations(self, maturity_data):
        """Generate specific recommendations based on KPI scores"""
        recommendations = []
        
        tag_score = maturity_data.get('tag_score', 0)
        budget_score = maturity_data.get('budget_score', 0)
        savings_score = maturity_data.get('savings_score', 0)
        utilization_score = maturity_data.get('utilization_score', 0)
        
        if tag_score < 3:
            recommendations.append("Implement comprehensive resource tagging strategy with mandatory cost center and environment tags")
        
        if budget_score < 3:
            recommendations.append("Establish proactive budget monitoring with automated alerts and forecasting")
        
        if savings_score < 3:
            recommendations.append("Execute immediate cost optimization initiatives focusing on idle resources and rightsizing")
        
        if utilization_score < 3:
            recommendations.append("Conduct resource utilization audit and implement automated scaling policies")
        
        # Always add general recommendations
        recommendations.extend([
            "Establish regular FinOps review meetings with stakeholders",
            "Implement cost anomaly detection and automated alerting",
            "Create cost optimization dashboards for continuous monitoring",
            "Develop cloud cost governance policies and procedures"
        ])
        
        return recommendations
    
    def generate_docx(
        self,
        cost_data: Dict,
        ri_data: Dict,
        recommendations: List[Dict],
        output_path: str
    ):
        """Generate DOCX report"""
        doc = Document()
        
        # Title with CloudXcelAI Branding
        title = doc.add_heading('CLOUDXCELAI FINOPTIMIZER', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.runs[0].font.color.rgb = RGBColor(0, 102, 204)
        
        subtitle = doc.add_paragraph('Cloud Financial Operations & Optimization Report')
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.runs[0].font.size = Pt(14)
        subtitle.runs[0].font.color.rgb = RGBColor(0, 61, 122)
        
        doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph()
        
        # Executive Summary
        doc.add_heading('Executive Summary', 1)
        
        summary_table = doc.add_table(rows=5, cols=2)
        summary_table.style = 'Light Grid Accent 1'
        
        summary_data = [
            ('Metric', 'Value'),
            ('Total Cost (6 months)', f"${cost_data['total_cost']:,.2f}"),
            ('Monthly Average', f"${cost_data['total_cost']/6:,.2f}"),
            ('Total Recommendations', str(len(recommendations))),
            ('Critical/High Priority', str(sum(1 for r in recommendations if r.get('severity') in ['Critical', 'High'])))
        ]
        
        for row_idx, (label, value) in enumerate(summary_data):
            row = summary_table.rows[row_idx]
            row.cells[0].text = label
            row.cells[1].text = value
            
            if row_idx == 0:
                for cell in row.cells:
                    cell.paragraphs[0].runs[0].font.bold = True
        
        doc.add_paragraph()
        
        # RI Analysis
        doc.add_heading('Reserved Instance Analysis', 1)
        ri_recommendations = ri_data['recommendations'].get('Recommendations', [])
        doc.add_paragraph(f"Active RIs: {len(ri_data['current_ris'])}")
        doc.add_paragraph(f"RI Purchase Recommendations: {len(ri_recommendations)}")
        doc.add_paragraph()
        
        # Recommendations
        doc.add_page_break()
        doc.add_heading('Cost Optimization Recommendations', 1)
        
        for idx, rec in enumerate(recommendations, 1):
            # Recommendation title
            doc.add_heading(f"{idx}. {rec.get('title', 'Recommendation')}", 2)
            
            # Details table
            details_table = doc.add_table(rows=4, cols=2)
            details_table.style = 'Light List Accent 1'
            
            details_data = [
                ('Severity', rec.get('severity', 'N/A')),
                ('Timeline', rec.get('timeline', 'N/A')),
                ('Estimated Savings', rec.get('estimated_savings', 'N/A')),
                ('Impact', rec.get('impact', 'N/A'))
            ]
            
            for row_idx, (label, value) in enumerate(details_data):
                row = details_table.rows[row_idx]
                row.cells[0].text = label
                row.cells[1].text = value
                row.cells[0].paragraphs[0].runs[0].font.bold = True
                
                # Color code severity
                if label == 'Severity':
                    severity = rec.get('severity', 'N/A')
                    if severity == 'Critical':
                        row.cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 0, 0)
                    elif severity == 'High':
                        row.cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 165, 0)
            
            doc.add_paragraph()
            
            # Description
            desc_para = doc.add_paragraph()
            desc_para.add_run('Description: ').bold = True
            desc_para.add_run(rec.get('description', 'N/A'))
            
            # Implementation steps
            if 'steps' in rec and rec['steps']:
                steps_para = doc.add_paragraph()
                steps_para.add_run('Implementation Steps:').bold = True
                
                for step in rec['steps']:
                    doc.add_paragraph(step, style='List Bullet')
            
            doc.add_paragraph()
        
        # Save document
        doc.save(output_path)
